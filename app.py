import io
import re
import datetime
import hashlib
import tempfile
import os
import smtplib
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
import streamlit as st
import pandas as pd
import polars as pl
from fpdf import FPDF
from recon_engine import execute_high_volume_reconciliation
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ==============================================================================
# --- Page Configuration & Executive Styling ---
# ==============================================================================
st.set_page_config(
    page_title="AutoReconFlow — Enterprise Financial Reconciliation, Tax & Audit Engine",
    page_icon="⚖️",
    layout="wide"
)

st.markdown("""
<style>
    @import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&display=swap');
    html, body, [class*="css"] {
        font-family: 'Inter', -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif;
        color: #1e293b;
    }
    .brand-hero {
        background: linear-gradient(135deg, #0A2540 0%, #1B365D 60%, #204066 100%);
        color: #ffffff;
        padding: 24px 30px;
        border-radius: 10px;
        margin-bottom: 20px;
        box-shadow: 0 4px 14px rgba(10, 37, 64, 0.15);
    }
    .brand-hero h1 {
        color: #ffffff !important;
        font-weight: 700;
        font-size: 25px;
        margin: 0 0 6px 0;
        letter-spacing: -0.5px;
    }
    .brand-hero p {
        color: #94a3b8 !important;
        font-size: 14px;
        margin: 0;
    }
    .compliance-badge-card {
        background: #f1f5f9;
        border: 1px solid #cbd5e1;
        border-left: 4px solid #1B365D;
        border-radius: 6px;
        padding: 12px 18px;
        margin-bottom: 22px;
        font-size: 13px;
        line-height: 1.5;
        color: #334155;
    }
    div[data-testid="stMetric"] {
        background-color: #ffffff;
        border: 1px solid #e2e8f0;
        border-radius: 8px;
        padding: 14px 18px;
        box-shadow: 0 1px 3px rgba(0, 0, 0, 0.05);
        border-left: 4px solid #1B365D;
    }
    div[data-testid="stMetricLabel"] p {
        font-size: 12.5px;
        font-weight: 600;
        color: #64748b;
        text-transform: uppercase;
        letter-spacing: 0.5px;
    }
    div[data-testid="stMetricValue"] div {
        color: #0f172a;
        font-weight: 700;
        font-size: 20px;
    }
    section[data-testid="stSidebar"] {
        background-color: #f8fafc;
        border-right: 1px solid #e2e8f0;
    }
    button[kind="primary"] {
        background: linear-gradient(180deg, #1B365D 0%, #0A2540 100%) !important;
        border: none !important;
        color: #ffffff !important;
        font-weight: 600 !important;
        padding: 8px 18px !important;
        border-radius: 6px !important;
        box-shadow: 0 2px 6px rgba(10, 37, 64, 0.2) !important;
        transition: all 0.2s ease;
    }
    button[kind="primary"]:hover {
        background: linear-gradient(180deg, #23426e 0%, #102d4d 100%) !important;
        box-shadow: 0 4px 12px rgba(10, 37, 64, 0.3) !important;
        transform: translateY(-1px);
    }
    div[data-testid="stDataFrame"] {
        border-radius: 8px;
        border: 1px solid #e2e8f0;
        overflow: hidden;
    }
</style>
""", unsafe_allow_html=True)

# ==============================================================================
# --- License Verification Engine (Deterministic SHA-256) ---
# ==============================================================================
def verify_license_key(email: str, license_key: str):
    if not email or not license_key:
        return False, None
    email_clean = email.strip().lower()
    key_clean = license_key.strip().upper()
    salt = "AutoReconFlowSecretKey2026"
    
    tier_map = {
        "ARF-PRAC": "Practitioner",
        "ARF-BIZ": "Business & Tax Suite",
        "ARF-ENT": "Enterprise Turnkey"
    }
    parts = key_clean.split("-")
    if len(parts) != 5:
        return False, None
    prefix = f"{parts[0]}-{parts[1]}"
    if prefix not in tier_map:
        return False, None
        
    tier_name = tier_map[prefix]
    current_date = datetime.datetime.utcnow()
    prev_month_date = current_date.replace(day=1) - datetime.timedelta(days=1)
    
    for m in [current_date.strftime("%Y%m"), prev_month_date.strftime("%Y%m")]:
        raw_str = f"{email_clean}:{tier_name}:{salt}:{m}"
        h = hashlib.sha256(raw_str.encode()).hexdigest().upper()
        if key_clean == f"{prefix}-{h[:4]}-{h[4:8]}-{h[8:12]}":
            return True, tier_name
    return False, None

# ==============================================================================
# --- Ingestion & String Sanitization Utilities ---
# ==============================================================================
def normalize_invoice_no(val):
    if val is None or pd.isna(val):
        return ""
    s = str(val).strip().upper()
    s = re.sub(r'[^A-Z0-9]', '', s)
    s = s.lstrip('0')
    return s if s else "0"

def clean_tax_id(val):
    if val is None or pd.isna(val):
        return ""
    return str(val).strip().upper().replace(" ", "").replace("-", "")

def robust_read_file(uploaded_file) -> pd.DataFrame:
    filename = uploaded_file.name.lower()
    uploaded_file.seek(0)
    
    if filename.endswith(".csv"):
        raw_bytes = uploaded_file.read()
        text = None
        for enc in ["utf-8-sig", "utf-8", "utf-16", "latin1"]:
            try:
                text = raw_bytes.decode(enc)
                if "\n" in text or "\r" in text:
                    break
            except Exception:
                continue
        if text is None:
            text = str(raw_bytes)
            
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        if not lines:
            return pd.DataFrame()
            
        delim = ","
        if "\t" in lines[0]:
            delim = "\t"
        elif ";" in lines[0]:
            delim = ";"
        elif "|" in lines[0]:
            delim = "|"
            
        try:
            df = pd.read_csv(io.StringIO(text), sep=delim, engine="python")
        except Exception:
            df = pd.read_csv(io.StringIO(text))
            
        if len(df.columns) == 1 and ("," in str(df.columns[0]) or ";" in str(df.columns[0])):
            actual_delim = "," if "," in str(df.columns[0]) else ";"
            df = pd.read_csv(io.StringIO(text), sep=actual_delim)
            
        df.columns = [str(c).strip().strip('"').strip("'").replace("\ufeff", "") for c in df.columns]
        return df.reset_index(drop=True)
        
    elif filename.endswith((".xlsx", ".xls")):
        df = pd.read_excel(uploaded_file)
        df.columns = [str(c).strip() for c in df.columns]
        return df.reset_index(drop=True)
        
    elif filename.endswith(".pdf"):
        import pdfplumber
        rows = []
        try:
            with pdfplumber.open(uploaded_file) as pdf:
                for page in pdf.pages:
                    tables = page.extract_tables()
                    for table in tables:
                        for r in table:
                            clean_r = [str(c).strip() if c is not None else "" for c in r]
                            if any(clean_r):
                                rows.append(clean_r)
        except Exception as e:
            st.error(f"Failed to read PDF: {str(e)}.")
            return pd.DataFrame()

        if not rows:
            st.warning("⚠️ No tabular data identified in this PDF.")
            return pd.DataFrame()

        headers = [re.sub(r'[^a-zA-Z0-9_]', '', h.replace(" ", "_").lower()) or f"col_{i}" for i, h in enumerate(rows[0])]
        df = pd.DataFrame(rows[1:], columns=headers)
        return df.reset_index(drop=True)
        
    return pd.DataFrame()

# ==============================================================================
# --- Pure Python Zero-Token Email Dispatcher ---
# ==============================================================================
def send_workflow_email(smtp_settings, recipient_email, subject, plain_text, html_content=None):
    if not smtp_settings.get("enabled", False):
        return True, "Simulation mode active (SMTP disabled)."
    
    server_host = smtp_settings.get("host", "smtp.gmail.com").strip()
    server_port = int(smtp_settings.get("port", 465))
    sender_user = smtp_settings.get("user", "").strip()
    sender_pass = smtp_settings.get("password", "").strip()

    if not sender_user or not sender_pass:
        return False, "SMTP sender username or password is missing."
    if not recipient_email or "@" not in recipient_email:
        return False, f"Invalid recipient email address: '{recipient_email}'"

    try:
        msg = MIMEMultipart("alternative")
        msg['From'] = f"AutoReconFlow Audit <{sender_user}>"
        msg['To'] = recipient_email
        msg['Subject'] = subject
        msg.attach(MIMEText(plain_text, 'plain'))
        if html_content:
            msg.attach(MIMEText(html_content, 'html'))

        if server_port == 465:
            with smtplib.SMTP_SSL(server_host, server_port, timeout=12) as server:
                server.login(sender_user, sender_pass)
                server.send_message(msg)
        else:
            with smtplib.SMTP(server_host, server_port, timeout=12) as server:
                server.starttls()
                server.login(sender_user, sender_pass)
                server.send_message(msg)
        return True, "Email successfully delivered."
    except Exception as e:
        return False, f"SMTP Dispatch Error: {str(e)}"

def build_html_email(title, header_desc, details_dict, action_button_url="https://www.autoreconflow.com"):
    rows_html = "".join([
        f"<tr><td style='padding:9px;font-weight:600;color:#1B365D;width:35%;border-bottom:1px solid #E2E8F0;'>{k}</td>"
        f"<td style='padding:9px;color:#334155;border-bottom:1px solid #E2E8F0;'>{v}</td></tr>"
        for k, v in details_dict.items()
    ])
    return f"""
    <html>
      <body style="font-family:'Segoe UI',Helvetica,Arial,sans-serif;margin:0;padding:24px;background-color:#F8FAFC;">
        <table align="center" border="0" cellpadding="0" cellspacing="0" width="600" style="background:#ffffff;border:1px solid #CBD5E1;border-radius:8px;overflow:hidden;box-shadow:0 4px 12px rgba(0,0,0,0.06);">
          <tr>
            <td bgcolor="#1B365D" style="padding:20px 28px;color:#ffffff;">
              <h2 style="margin:0;font-size:19px;font-weight:700;">⚖️ AutoReconFlow | Financial Governance Notice</h2>
            </td>
          </tr>
          <tr>
            <td style="padding:28px;">
              <h3 style="margin-top:0;color:#0F172A;font-size:17px;">{title}</h3>
              <p style="color:#475569;font-size:14px;line-height:1.6;">{header_desc}</p>
              <table width="100%" cellpadding="0" cellspacing="0" style="margin:20px 0;border-collapse:collapse;font-size:13px;">
                {rows_html}
              </table>
              <div style="text-align:center;margin-top:28px;">
                <a href="{action_button_url}" style="background-color:#1B365D;color:#ffffff;text-decoration:none;padding:12px 26px;border-radius:5px;font-weight:600;display:inline-block;box-shadow:0 2px 6px rgba(27,54,93,0.3);">Open Workpaper & Verify Record</a>
              </div>
            </td>
          </tr>
          <tr>
            <td bgcolor="#F1F5F9" style="padding:14px 28px;font-size:11px;color:#64748B;text-align:center;">
              SOX 404 / ISA-500 Four-Eyes Governance Automation | Tamper-Evident SHA-256 Audit Trails
            </td>
          </tr>
        </table>
      </body>
    </html>
    """

def build_vendor_gst_notice(buyer_name, buyer_gstin, supplier_name, missing_invoices_df):
    rows = ""
    total_blocked_tax = 0.0
    for _, r in missing_invoices_df.iterrows():
        inv_no = r.get('_inv_PR', r.get('Invoice_Number', 'N/A'))
        dt = r.get('Invoice_Date', 'N/A')
        taxable = float(r.get('Taxable_Value', 0.0))
        tax = float(r.get('Risk_Amount_ITC', r.get('_pr_tax_total', 0.0)))
        total_blocked_tax += tax
        rows += f"""
        <tr>
            <td style='padding:8px;border:1px solid #CBD5E1;'>{inv_no}</td>
            <td style='padding:8px;border:1px solid #CBD5E1;'>{dt}</td>
            <td style='padding:8px;border:1px solid #CBD5E1;text-align:right;'>₹{taxable:,.2f}</td>
            <td style='padding:8px;border:1px solid #CBD5E1;text-align:right;font-weight:bold;color:#b91c1c;'>₹{tax:,.2f}</td>
        </tr>
        """
    html = f"""
    <html>
      <body style="font-family:'Segoe UI',sans-serif;color:#1e293b;padding:20px;background:#f8fafc;">
        <div style="background:#fff;border:1px solid #cbd5e1;border-radius:8px;padding:24px;max-width:650px;margin:auto;">
          <h3 style="color:#1b365d;margin-top:0;">Statutory Notice: Unfiled Invoices in GSTR-2B [Section 16(2)(aa)]</h3>
          <p>Dear Accounts Team at <strong>{supplier_name}</strong>,</p>
          <p>During our monthly GST reconciliation for GSTR-3B filing, the following invoices recorded in our purchase books <strong>have not been reflected in our auto-drafted GSTR-2B</strong> statement:</p>
          <table width="100%" style="border-collapse:collapse;font-size:13px;margin:16px 0;">
            <thead>
              <tr style="background:#1b365d;color:#fff;">
                <th style="padding:8px;border:1px solid #1b365d;">Invoice No</th>
                <th style="padding:8px;border:1px solid #1b365d;">Date</th>
                <th style="padding:8px;border:1px solid #1b365d;text-align:right;">Taxable Value</th>
                <th style="padding:8px;border:1px solid #1b365d;text-align:right;">Blocked Tax</th>
              </tr>
            </thead>
            <tbody>{rows}</tbody>
            <tfoot>
              <tr style="background:#f1f5f9;font-weight:bold;">
                <td colspan="3" style="padding:8px;border:1px solid #cbd5e1;text-align:right;">Total Blocked ITC:</td>
                <td style="padding:8px;border:1px solid #cbd5e1;text-align:right;color:#b91c1c;">₹{total_blocked_tax:,.2f}</td>
              </tr>
            </tfoot>
          </table>
          <p style="font-size:13px;line-height:1.5;">Under <strong>Section 16(2)(aa) of the CGST Act</strong>, we cannot legally claim this Input Tax Credit until these invoices are reported in your GSTR-1/IFF. Please ensure immediate filing to prevent payment withholding or interest debit notes under Section 50.</p>
          <hr style="border:none;border-top:1px solid #e2e8f0;margin:20px 0;" />
          <p style="font-size:12px;color:#64748b;margin:0;">Issued by: <strong>{buyer_name}</strong> | Identification: {buyer_gstin}</p>
        </div>
      </body>
    </html>
    """
    return html, total_blocked_tax

# ==============================================================================
# --- PDF Executive Sign-Off Certificate Generator ---
# ==============================================================================
def sanitize_pdf_text(text: str) -> str:
    if not text:
        return ""
    replacements = {"—": "-", "–": "-", "’": "'", "‘": "'", "“": '"', "”": '"', "•": "*", "…": "...", "±": "+/-"}
    for orig, rep in replacements.items():
        text = text.replace(orig, rep)
    return text.encode("latin-1", "replace").decode("latin-1")

class AuditCertificatePDF(FPDF):
    def __init__(self, logo_path=None):
        super().__init__()
        self.logo_path = logo_path

    def header(self):
        self.set_fill_color(27, 54, 93)
        self.rect(0, 0, 210, 16, 'F')
        self.set_text_color(255, 255, 255)
        self.set_font('Helvetica', 'B', 11)
        self.cell(0, 4, 'AUTORECONFLOW - FINANCIAL RECONCILIATION, TAX & AUDIT CERTIFICATE', 0, 1, 'C')
        self.ln(6)

    def footer(self):
        self.set_y(-18)
        self.set_font('Helvetica', 'I', 8)
        self.set_text_color(128, 128, 128)
        self.cell(0, 4, 'Confidential - Statutory Audit Evidence | PCAOB AS 2201 / SOX 404 / ISA-500 Standard', 0, 1, 'C')
        self.cell(0, 4, f'Page {self.page_no()}', 0, 0, 'C')

def generate_pdf_signoff_certificate(matched_count, fuzzy_count, un_a_count, un_b_count, signoff_data, tolerance, logo_bytes=None, company_info=None, query_trail="") -> bytes:
    logo_temp_file = None
    if logo_bytes:
        try:
            logo_temp = tempfile.NamedTemporaryFile(delete=False, suffix=".png")
            logo_temp.write(logo_bytes)
            logo_temp.flush()
            logo_temp.close()
            logo_temp_file = logo_temp.name
        except Exception:
            logo_temp_file = None

    pdf = AuditCertificatePDF(logo_path=logo_temp_file)
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)

    comp_name = sanitize_pdf_text(company_info.get("name", "AutoReconFlow Enterprise"))
    comp_addr = sanitize_pdf_text(company_info.get("address", ""))
    comp_uid = sanitize_pdf_text(company_info.get("uid", ""))

    y_brand_start = pdf.get_y()
    pdf.set_xy(10, y_brand_start)
    pdf.set_font('Helvetica', 'B', 13)
    pdf.set_text_color(27, 54, 93)
    pdf.cell(120, 6, comp_name, ln=True)

    if comp_uid:
        pdf.set_font('Helvetica', 'B', 8)
        pdf.set_text_color(70, 70, 70)
        pdf.cell(120, 4.5, f"Legal / Tax Identification: {comp_uid}", ln=True)

    if comp_addr:
        pdf.set_font('Helvetica', '', 7.5)
        pdf.set_text_color(100, 100, 100)
        pdf.multi_cell(125, 3.8, f"Registered Office: {comp_addr}")

    if logo_temp_file and os.path.exists(logo_temp_file):
        try:
            pdf.image(logo_temp_file, x=150, y=y_brand_start - 2, h=16)
        except Exception:
            pass

    pdf.set_y(max(pdf.get_y() + 3, y_brand_start + 18))
    pdf.set_draw_color(203, 213, 225)
    pdf.line(10, pdf.get_y(), 200, pdf.get_y())
    pdf.ln(4)

    pdf.set_font('Helvetica', 'B', 14)
    pdf.set_text_color(27, 54, 93)
    pdf.cell(0, 7, 'Financial Reconciliation & Tax Audit Certificate', ln=True, align='L')
    pdf.set_font('Helvetica', 'I', 8.5)
    pdf.set_text_color(100, 100, 100)
    pdf.cell(0, 4.5, 'Compliance Certification under PCAOB AS 2201 / SOX Section 404 & ISA-500 Standards', ln=True, align='L')
    pdf.ln(3)

    pdf.set_fill_color(248, 250, 252)
    pdf.set_draw_color(203, 213, 225)
    pdf.rect(10, pdf.get_y(), 190, 22, 'DF')
    pdf.set_xy(14, pdf.get_y() + 2)

    pdf.set_font('Helvetica', 'B', 8.5)
    pdf.set_text_color(50, 50, 50)
    pdf.cell(45, 4.5, 'Execution Timestamp (UTC):', 0, 0)
    pdf.set_font('Helvetica', '', 8.5)
    pdf.cell(55, 4.5, sanitize_pdf_text(signoff_data.get('preparer_date', '')), 0, 0)

    pdf.set_font('Helvetica', 'B', 8.5)
    pdf.cell(40, 4.5, 'Workpaper Reference:', 0, 0)
    pdf.set_font('Helvetica', '', 8.5)
    pdf.cell(45, 4.5, sanitize_pdf_text(signoff_data.get('wp_ref', '')), 0, 1)

    pdf.set_x(14)
    pdf.set_font('Helvetica', 'B', 8.5)
    pdf.cell(45, 4.5, 'Variance Tolerance:', 0, 0)
    pdf.set_font('Helvetica', '', 8.5)
    pdf.cell(55, 4.5, f'+/- {tolerance:.2f} Currency Units', 0, 0)

    pdf.set_font('Helvetica', 'B', 8.5)
    pdf.cell(40, 4.5, 'Supervisory Verdict:', 0, 0)
    verdict = sanitize_pdf_text(signoff_data.get('review_verdict', 'Approved / Cleared'))
    pdf.set_font('Helvetica', 'B', 8.5)
    pdf.set_text_color(16, 124, 65) if 'Approved' in verdict else pdf.set_text_color(180, 40, 40)
    pdf.cell(45, 4.5, verdict, 0, 1)

    pdf.ln(6)

    pdf.set_text_color(255, 255, 255)
    pdf.set_fill_color(27, 54, 93)
    pdf.set_font('Helvetica', 'B', 8.5)
    pdf.cell(110, 6.5, 'Reconciliation Assertion Category', 1, 0, 'L', fill=True)
    pdf.cell(35, 6.5, 'Record Count', 1, 0, 'C', fill=True)
    pdf.cell(45, 6.5, 'Audit Status', 1, 1, 'C', fill=True)

    table_data = [
        ("Tier-1 Deterministic Matches (Exact Key & Net Match)", str(matched_count), "Cleared"),
        ("Tier-2 Discrepancy & Threshold Review (Materiality Tracking)", str(fuzzy_count), "Supervisory Review" if fuzzy_count > 0 else "Cleared"),
        ("Source A Exceptions (Subledger Missing in External Records)", str(un_a_count), "Remediation Req." if un_a_count > 0 else "Cleared"),
        ("Source B Exceptions (External Records Missing in Subledger)", str(un_b_count), "Remediation Req." if un_b_count > 0 else "Cleared"),
    ]

    pdf.set_text_color(40, 40, 40)
    pdf.set_font('Helvetica', '', 8)
    for cat, count, status in table_data:
        pdf.cell(110, 6, f"  {cat}", 1, 0, 'L')
        pdf.cell(35, 6, count, 1, 0, 'C')
        pdf.set_text_color(16, 124, 65) if status == "Cleared" else pdf.set_text_color(180, 40, 40)
        pdf.cell(45, 6, status, 1, 1, 'C')
        pdf.set_text_color(40, 40, 40)

    pdf.ln(7)

    pdf.set_font('Helvetica', 'B', 10.5)
    pdf.set_text_color(27, 54, 93)
    pdf.cell(0, 5.5, 'Four-Eyes Internal Control Attestation', ln=True)
    pdf.ln(2)

    y_start = pdf.get_y()
    pdf.set_fill_color(248, 250, 252)
    pdf.rect(10, y_start, 92, 40, 'DF')
    pdf.set_xy(14, y_start + 2.5)
    pdf.set_font('Helvetica', 'B', 8.5)
    pdf.set_text_color(27, 54, 93)
    pdf.cell(85, 4.5, 'PREPARER (First Line of Defense)', 0, 1)

    pdf.set_x(14)
    pdf.set_font('Helvetica', '', 8)
    pdf.set_text_color(50, 50, 50)
    pdf.cell(85, 4.2, f"Name: {sanitize_pdf_text(signoff_data.get('preparer_name'))}", 0, 1)
    pdf.set_x(14)
    pdf.cell(85, 4.2, f"System / Emp ID: {sanitize_pdf_text(signoff_data.get('preparer_id'))}", 0, 1)
    pdf.set_x(14)
    pdf.cell(85, 4.2, f"Timestamp: {sanitize_pdf_text(signoff_data.get('preparer_date'))}", 0, 1)
    sig_text = "[VERIFIED ELECTRONIC SIGNATURE - COMPLETED]" if signoff_data.get('preparer_signed') else "[UNSIGNED]"
    pdf.set_font('Helvetica', 'B', 7.5)
    pdf.set_text_color(16, 124, 65) if signoff_data.get('preparer_signed') else pdf.set_text_color(180, 40, 40)
    pdf.cell(85, 5.5, sig_text, 0, 1)

    pdf.set_fill_color(248, 250, 252)
    pdf.rect(108, y_start, 92, 40, 'DF')
    pdf.set_xy(112, y_start + 2.5)
    pdf.set_font('Helvetica', 'B', 8.5)
    pdf.set_text_color(27, 54, 93)
    pdf.cell(85, 4.5, 'REVIEWER (Supervisory Review / Oversight)', 0, 1)

    pdf.set_x(112)
    pdf.set_font('Helvetica', '', 8)
    pdf.set_text_color(50, 50, 50)
    pdf.cell(85, 4.2, f"Name: {sanitize_pdf_text(signoff_data.get('reviewer_name'))}", 0, 1)
    pdf.set_x(112)
    pdf.cell(85, 4.2, f"Designation: {sanitize_pdf_text(signoff_data.get('reviewer_title'))}", 0, 1)
    pdf.set_x(112)
    pdf.cell(85, 4.2, f"Timestamp: {sanitize_pdf_text(signoff_data.get('reviewer_date'))}", 0, 1)
    pdf.set_x(112)
    pdf.set_font('Helvetica', 'B', 7.5)
    pdf.set_text_color(16, 124, 65)
    pdf.cell(85, 5.5, f"[SUPERVISORY APPROVAL: {verdict.upper()}]", 0, 1)

    pdf.set_y(y_start + 44)

    pdf.set_font('Helvetica', 'B', 8.5)
    pdf.set_text_color(27, 54, 93)
    pdf.cell(0, 4.5, 'Supervisory Audit Notes & Resolution Trail:', ln=True)
    pdf.set_font('Helvetica', 'I', 8)
    pdf.set_text_color(60, 60, 60)
    notes_clean = sanitize_pdf_text(signoff_data.get('audit_comments', 'No comments noted.') or 'None noted.')
    if query_trail:
        notes_clean += f" | {sanitize_pdf_text(query_trail)}"
    pdf.multi_cell(0, 4, notes_clean)
    pdf.ln(3)

    raw_fingerprint = f"{comp_uid}_{signoff_data.get('wp_ref')}_{matched_count}_{fuzzy_count}_{un_a_count}_{un_b_count}_{signoff_data.get('preparer_date')}_{signoff_data.get('reviewer_date')}_{tolerance}_{query_trail}"
    doc_hash = hashlib.sha256(raw_fingerprint.encode()).hexdigest().upper()

    pdf.set_fill_color(240, 240, 240)
    pdf.set_font('Courier', 'B', 7.5)
    pdf.set_text_color(70, 70, 70)
    pdf.cell(0, 6, f" SHA-256 DIGITAL IMMUTABILITY AUDIT HASH: {doc_hash} ", 1, 1, 'C', fill=True)

    result_bytes = bytes(pdf.output())
    if logo_temp_file and os.path.exists(logo_temp_file):
        try:
            os.remove(logo_temp_file)
        except Exception:
            pass

    return result_bytes

# ==============================================================================
# --- 1. GST ITC Reconciler Logic ---
# ==============================================================================
def execute_gst_itc_reconciliation(df_pr: pd.DataFrame, df_2b: pd.DataFrame, map_pr: dict, map_2b: dict, tolerance=2.0):
    pr = df_pr.copy()
    g2b = df_2b.copy()

    pr['_gstin'] = pr[map_pr['gstin']].apply(clean_tax_id)
    pr['_inv'] = pr[map_pr['inv']].apply(normalize_invoice_no)
    pr['_key'] = pr['_gstin'] + "_" + pr['_inv']

    g2b['_gstin'] = g2b[map_2b['gstin']].apply(clean_tax_id)
    g2b['_inv'] = g2b[map_2b['inv']].apply(normalize_invoice_no)
    g2b['_key'] = g2b['_gstin'] + "_" + g2b['_inv']

    for f in ['taxable', 'igst', 'cgst', 'sgst']:
        col_pr = map_pr.get(f)
        col_2b = map_2b.get(f)
        pr[f'_pr_{f}'] = pd.to_numeric(pr[col_pr], errors='coerce').fillna(0.0) if col_pr else 0.0
        g2b[f'_2b_{f}'] = pd.to_numeric(g2b[col_2b], errors='coerce').fillna(0.0) if col_2b else 0.0

    pr['_pr_tax_total'] = pr['_pr_igst'] + pr['_pr_cgst'] + pr['_pr_sgst']
    g2b['_2b_tax_total'] = g2b['_2b_igst'] + g2b['_2b_cgst'] + g2b['_2b_sgst']

    merged = pd.merge(pr, g2b, on='_key', how='outer', suffixes=('_PR', '_2B'), indicator=True)

    b1_matched, b2_missing_in_books, b3_missing_in_2b, b4_variances = [], [], [], []

    for idx, row in merged.iterrows():
        status = row['_merge']
        if status == 'both':
            tax_diff = abs(row['_pr_tax_total'] - row['_2b_tax_total'])
            taxable_diff = abs(row['_pr_taxable'] - row['_2b_taxable'])
            
            if tax_diff <= tolerance and taxable_diff <= tolerance:
                row_dict = row.to_dict()
                row_dict['Audit_Finding'] = "Eligible ITC - Direct Match [Section 16(2)(aa) Cleared]"
                row_dict['Claimable_ITC_Total'] = row['_2b_tax_total']
                b1_matched.append(row_dict)
            else:
                row_dict = row.to_dict()
                reasons = []
                if tax_diff > tolerance:
                    reasons.append(f"Tax Head Diff: ₹{tax_diff:.2f}")
                if taxable_diff > tolerance:
                    reasons.append(f"Taxable Value Diff: ₹{taxable_diff:.2f}")
                row_dict['Audit_Finding'] = f"Value Discrepancy: {', '.join(reasons)}"
                row_dict['PR_Total_Tax'] = row['_pr_tax_total']
                row_dict['2B_Total_Tax'] = row['_2b_tax_total']
                row_dict['Tax_Variance'] = row['_pr_tax_total'] - row['_2b_tax_total']
                b4_variances.append(row_dict)
        elif status == 'left_only':
            row_dict = row.to_dict()
            row_dict['Audit_Finding'] = "BLOCKED ITC: In Books, Unfiled in 2B [Section 16(2)(aa) Ineligible]"
            row_dict['Risk_Amount_ITC'] = row['_pr_tax_total']
            b3_missing_in_2b.append(row_dict)
        elif status == 'right_only':
            row_dict = row.to_dict()
            row_dict['Audit_Finding'] = "UNRECORDED: In 2B, Not in Books [Unclaimed Ledger Credit]"
            row_dict['Available_Unclaimed_ITC'] = row['_2b_tax_total']
            b2_missing_in_books.append(row_dict)

    df_b1 = pd.DataFrame(b1_matched).drop(columns=['_key', '_merge', '_gstin_PR', '_gstin_2B', '_inv_PR', '_inv_2B'], errors='ignore')
    df_b2 = pd.DataFrame(b2_missing_in_books).drop(columns=['_key', '_merge', '_gstin_PR', '_gstin_2B', '_inv_PR', '_inv_2B'], errors='ignore')
    df_b3 = pd.DataFrame(b3_missing_in_2b).drop(columns=['_key', '_merge', '_gstin_PR', '_gstin_2B', '_inv_PR', '_inv_2B'], errors='ignore')
    df_b4 = pd.DataFrame(b4_variances).drop(columns=['_key', '_merge', '_gstin_PR', '_gstin_2B', '_inv_PR', '_inv_2B'], errors='ignore')

    return df_b1, df_b2, df_b3, df_b4

def generate_gst_itc_excel_workpaper(df_b1, df_b2, df_b3, df_b4, company_name, client_gstin) -> bytes:
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        workbook = writer.book
        ws_sum = workbook.create_sheet(title="ITC Audit Summary", index=0)
        ws_sum.views.sheetView[0].showGridLines = True
        
        navy_header = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
        white_bold = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        bold_f = Font(name="Calibri", size=11, bold=True)
        reg_f = Font(name="Calibri", size=11)
        
        ws_sum["A1"] = f"{company_name} — GSTR-2B vs Purchase Register Audit Workpaper"
        ws_sum["A1"].font = Font(name="Calibri", size=14, bold=True, color="1B365D")
        ws_sum["A2"] = f"Statutory Compliance: CGST Section 16(2)(aa) & Rule 36(4) | Identification: {client_gstin} | Generated: {datetime.datetime.utcnow().strftime('%Y-%m-%d %H:%M:%S UTC')}"
        ws_sum["A2"].font = Font(name="Calibri", size=9.5, italic=True, color="595959")
        
        headers = ["GSTR-3B Statutory Bucket", "Invoice Count", "Total Tax / ITC Impact (₹)", "Statutory Recommendation & Action"]
        ws_sum.append([])
        ws_sum.append(headers)
        
        for c in range(1, 5):
            cell = ws_sum.cell(row=4, column=c)
            cell.fill = navy_header
            cell.font = white_bold
            
        b1_tax = float(df_b1['Claimable_ITC_Total'].sum()) if len(df_b1) > 0 and 'Claimable_ITC_Total' in df_b1.columns else 0.0
        b2_tax = float(df_b2['Available_Unclaimed_ITC'].sum()) if len(df_b2) > 0 and 'Available_Unclaimed_ITC' in df_b2.columns else 0.0
        b3_tax = float(df_b3['Risk_Amount_ITC'].sum()) if len(df_b3) > 0 and 'Risk_Amount_ITC' in df_b3.columns else 0.0
        b4_tax = float(df_b4['Tax_Variance'].sum()) if len(df_b4) > 0 and 'Tax_Variance' in df_b4.columns else 0.0
        
        rows = [
            ("Bucket 1: Direct Match (Eligible ITC)", len(df_b1), b1_tax, "Claim 100% in Table 4(A)(5) of GSTR-3B under Sec 16(2)(aa)"),
            ("Bucket 2: In 2B but Missing in Books", len(df_b2), b2_tax, "Verify unbooked expenses; pass purchase voucher before claiming"),
            ("Bucket 3: In Books but Missing in 2B", len(df_b3), b3_tax, "DO NOT CLAIM. Block credit; issue notice to non-filing supplier"),
            ("Bucket 4: Value & Tax Head Discrepancies", len(df_b4), b4_tax, "Rectify voucher tax heads or request supplier GSTR-1A amendment")
        ]
        
        for r_data in rows:
            ws_sum.append(list(r_data))
            
        for r in range(5, 9):
            ws_sum.cell(row=r, column=1).font = bold_f
            ws_sum.cell(row=r, column=2).font = reg_f
            ws_sum.cell(row=r, column=3).font = bold_f
            ws_sum.cell(row=r, column=3).number_format = "₹#,##0.00"
            ws_sum.cell(row=r, column=4).font = reg_f

        ws_sum.column_dimensions["A"].width = 44
        ws_sum.column_dimensions["B"].width = 16
        ws_sum.column_dimensions["C"].width = 28
        ws_sum.column_dimensions["D"].width = 65

        sheets = {
            "Bucket 1 - Eligible ITC": df_b1,
            "Bucket 2 - Missing in Books": df_b2,
            "Bucket 3 - Missing in 2B": df_b3,
            "Bucket 4 - Mismatches": df_b4
        }
        for title, df_data in sheets.items():
            clean_df = df_data.copy()
            clean_df.to_excel(writer, sheet_name=title, index=False)
            ws = writer.sheets[title]
            ws.views.sheetView[0].showGridLines = True
            for col_idx in range(1, len(clean_df.columns) + 1):
                c_cell = ws.cell(row=1, column=col_idx)
                c_cell.fill = navy_header
                c_cell.font = white_bold
            for col in ws.columns:
                max_l = max(len(str(cell.value or "")) for cell in col)
                col_let = get_column_letter(col[0].column)
                ws.column_dimensions[col_let].width = min(max(max_l + 3, 14), 45)

    return output.getvalue()

# ==============================================================================
# --- 2. Two-Sided International VAT Reverse Charge (RCM) Engine ---
# ==============================================================================
def execute_two_sided_vat_rcm_reconciliation(df_books: pd.DataFrame, df_ext: pd.DataFrame, map_books: dict, map_ext: dict, applicable_vat_rate=20.0, tolerance=1.0):
    bks = df_books.copy()
    ext = df_ext.copy()

    # Normalization keys
    bks['_tax_id'] = bks[map_books['vat_id']].apply(clean_tax_id)
    bks['_inv'] = bks[map_books['invoice']].apply(normalize_invoice_no)
    bks['_key'] = bks['_tax_id'] + "_" + bks['_inv']

    ext['_tax_id'] = ext[map_ext['vat_id']].apply(clean_tax_id)
    ext['_inv'] = ext[map_ext['invoice']].apply(normalize_invoice_no)
    ext['_key'] = ext['_tax_id'] + "_" + ext['_inv']

    # Numeric fields
    bks['_bks_net'] = pd.to_numeric(bks[map_books['net_amt']], errors='coerce').fillna(0.0)
    bks['_bks_self_tax'] = pd.to_numeric(bks[map_books['self_assessed']], errors='coerce').fillna(0.0) if map_books.get('self_assessed') else 0.0

    ext['_ext_net'] = pd.to_numeric(ext[map_ext['net_amt']], errors='coerce').fillna(0.0)
    ext['_ext_charged_tax'] = pd.to_numeric(ext[map_ext['charged_vat']], errors='coerce').fillna(0.0)

    # Perform full outer join
    merged = pd.merge(bks, ext, on='_key', how='outer', suffixes=('_BOOKS', '_EXT'), indicator=True)

    b1_matched, b2_missing_in_books, b3_missing_in_ext, b4_discrepancies = [], [], [], []

    for idx, row in merged.iterrows():
        status = row['_merge']
        if status == 'both':
            net_diff = abs(row['_bks_net'] - row['_ext_net'])
            expected_rcm = round(row['_bks_net'] * (applicable_vat_rate / 100.0), 2)
            charged_foreign_tax = row['_ext_charged_tax']
            self_assessed_tax = row['_bks_self_tax']
            is_rcm = str(row.get(map_books.get('is_rcm', ''), 'YES')).strip().upper() in ['YES', 'TRUE', '1', 'Y', 'RCM']
            
            if is_rcm and charged_foreign_tax > tolerance:
                r_dict = row.to_dict()
                r_dict['Audit_Finding'] = f"UNLAWFUL FOREIGN TAX: Supplier charged {charged_foreign_tax:.2f} foreign VAT on cross-border RCM service."
                r_dict['Compliance_Status'] = "UNLAWFUL_SUPPLIER_VAT"
                r_dict['Correction_Amount'] = charged_foreign_tax
                b4_discrepancies.append(r_dict)
            elif net_diff > tolerance:
                r_dict = row.to_dict()
                r_dict['Audit_Finding'] = f"NET BASE VARIANCE: Books ({row['_bks_net']:.2f}) vs Ext ({row['_ext_net']:.2f}) diff: {net_diff:.2f}"
                r_dict['Compliance_Status'] = "BASE_DISCREPANCY"
                r_dict['Correction_Amount'] = net_diff
                b4_discrepancies.append(r_dict)
            elif self_assessed_tax > 0 and abs(self_assessed_tax - expected_rcm) > tolerance:
                r_dict = row.to_dict()
                r_dict['Audit_Finding'] = f"SELF-ASSESSMENT RATE ERROR: Recorded deemed VAT ({self_assessed_tax:.2f}) vs Statutory ({expected_rcm:.2f})"
                r_dict['Compliance_Status'] = "RATE_MISMATCH"
                r_dict['Correction_Amount'] = expected_rcm - self_assessed_tax
                b4_discrepancies.append(r_dict)
            else:
                r_dict = row.to_dict()
                r_dict['Audit_Finding'] = f"STATUTORY CLEARED: Valid RCM match. Deemed output matched with input recovery ({expected_rcm:.2f})."
                r_dict['Compliance_Status'] = "CLEARED_RCM"
                r_dict['Deemed_VAT'] = expected_rcm
                b1_matched.append(r_dict)
        elif status == 'left_only':
            r_dict = row.to_dict()
            r_dict['Audit_Finding'] = "In Books but missing in External/Supplier Statement"
            r_dict['Compliance_Status'] = "MISSING_IN_EXTERNAL"
            b3_missing_in_ext.append(r_dict)
        elif status == 'right_only':
            r_dict = row.to_dict()
            r_dict['Audit_Finding'] = "Billed by Supplier / Reported in External, missing in Books"
            r_dict['Compliance_Status'] = "UNRECORDED_LIABILITY"
            b2_missing_in_books.append(r_dict)

    df_b1 = pd.DataFrame(b1_matched).drop(columns=['_key', '_merge', '_tax_id_BOOKS', '_tax_id_EXT', '_inv_BOOKS', '_inv_EXT'], errors='ignore')
    df_b2 = pd.DataFrame(b2_missing_in_books).drop(columns=['_key', '_merge', '_tax_id_BOOKS', '_tax_id_EXT', '_inv_BOOKS', '_inv_EXT'], errors='ignore')
    df_b3 = pd.DataFrame(b3_missing_in_ext).drop(columns=['_key', '_merge', '_tax_id_BOOKS', '_tax_id_EXT', '_inv_BOOKS', '_inv_EXT'], errors='ignore')
    df_b4 = pd.DataFrame(b4_discrepancies).drop(columns=['_key', '_merge', '_tax_id_BOOKS', '_tax_id_EXT', '_inv_BOOKS', '_inv_EXT'], errors='ignore')

    return df_b1, df_b2, df_b3, df_b4

# ==============================================================================
# --- Universal Multi-Tab Audit Workpaper Generator (openpyxl) ---
# ==============================================================================
def generate_consolidated_workbook(matched_df, fuzzy_df, un_a_df, un_b_df, signoff_data, tolerance=1.0, company_info=None, query_obj=None) -> bytes:
    output = io.BytesIO()
    timestamp_utc = datetime.datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S UTC")

    comp_name = company_info.get("name", "AutoReconFlow Enterprise") if company_info else "AutoReconFlow Enterprise"
    comp_addr = company_info.get("address", "") if company_info else ""
    comp_uid = company_info.get("uid", "") if company_info else ""

    navy_header_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    section_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    white_bold_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=11, bold=True)
    regular_font = Font(name="Calibri", size=11)
    italic_font = Font(name="Calibri", size=10, italic=True, color="595959")
    title_font = Font(name="Calibri", size=14, bold=True, color="1B365D")
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )

    m_copy = matched_df.copy()
    f_copy = fuzzy_df.copy()
    a_copy = un_a_df.copy()
    b_copy = un_b_df.copy()

    if query_obj and query_obj.get("status") == "APPROVED" and query_obj.get("item_id"):
        flag_id = str(query_obj["item_id"])
        trail_entry = f"[REVIEWER {query_obj['reviewer_query_time']}]: {query_obj['reviewer_query']} | [PREPARER {query_obj['preparer_reply_time']}]: {query_obj['preparer_reply']}"
        for df_target in [f_copy, a_copy, b_copy]:
            match_col = None
            for c in df_target.columns:
                if any(k in str(c).lower() for k in ["id", "invoice", "ref", "trans"]):
                    match_col = c
                    break
            if match_col:
                df_target["Audit_Clarification_Trail"] = df_target[match_col].apply(
                    lambda x: trail_entry if str(x) == flag_id else ""
                )

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        workbook = writer.book

        ws_summary = workbook.create_sheet(title="Audit & Control Summary", index=0)
        ws_summary.views.sheetView[0].showGridLines = True

        ws_summary["A1"] = f"{comp_name} — Financial Reconciliation, Tax & Control Workpaper"
        ws_summary["A1"].font = title_font
        ws_summary.merge_cells("A1:C1")

        ws_summary["A2"] = f"Legal Identification: {comp_uid} | Registered Office: {comp_addr}"
        ws_summary["A2"].font = italic_font
        ws_summary.merge_cells("A2:C2")

        ws_summary["A3"] = "SOX 404 / ISA-500 Standard Financial Statement & Tax Assertion Audit Trail"
        ws_summary["A3"].font = Font(name="Calibri", size=9, italic=True, color="64748B")
        ws_summary.merge_cells("A3:C3")

        metadata = [
            ("Execution Timestamp (UTC):", timestamp_utc),
            ("Tolerance Applied:", f"± {tolerance:.2f} Currency Units"),
            ("Matching Methodology:", "3-Tier (Deterministic Exact -> Levenshtein/Jaro-Winkler Fuzzy)"),
            ("Control Assertion Target:", "Completeness, Accuracy, and Valuation (Financial & Tax Controls)")
        ]
        ws_summary.append([])
        for label, val in metadata:
            ws_summary.append([label, val])

        for r in range(5, 9):
            ws_summary.cell(row=r, column=1).font = bold_font
            ws_summary.cell(row=r, column=2).font = regular_font

        ws_summary.append([])
        ws_summary.append(["Reconciliation Category", "Record Count", "Control Status"])
        summary_header_row = 10
        for col_idx in range(1, 4):
            cell = ws_summary.cell(row=summary_header_row, column=col_idx)
            cell.fill = navy_header_fill
            cell.font = white_bold_font
            cell.alignment = Alignment(horizontal="center" if col_idx > 1 else "left")

        summary_rows = [
            ("Tier-1 Deterministic Matches (Subledger to External)", len(m_copy), "Cleared"),
            ("Tier-2 Discrepancy & Threshold Review (Materiality Tracking)", len(f_copy), "Requires Supervisory Review" if len(f_copy) > 0 else "Cleared"),
            ("Source A Exceptions (Ledger Missing in External)", len(a_copy), "Action Required" if len(a_copy) > 0 else "Cleared"),
            ("Source B Exceptions (External Missing in Ledger)", len(b_copy), "Action Required" if len(b_copy) > 0 else "Cleared"),
        ]

        for item in summary_rows:
            ws_summary.append(list(item))

        for r in range(11, 15):
            for c in range(1, 4):
                cell = ws_summary.cell(row=r, column=c)
                cell.font = regular_font
                cell.border = thin_border
                if c == 2:
                    cell.alignment = Alignment(horizontal="right")
                    cell.number_format = "#,##0"
                elif c == 3:
                    cell.alignment = Alignment(horizontal="center")

        ws_summary.append([])
        sign_off_header_row = 16
        ws_summary.cell(row=sign_off_header_row, column=1, value="Internal Control Sign-Off (SOX 404 / Four-Eyes Governance)").font = bold_font
        ws_summary.cell(row=sign_off_header_row, column=1).fill = section_fill
        ws_summary.merge_cells(start_row=sign_off_header_row, start_column=1, end_row=sign_off_header_row, end_column=3)

        sig_status = "VERIFIED / ELECTRONICALLY SIGNED" if signoff_data.get("preparer_signed") else "PENDING SIGNATURE"
        signoff_entries = [
            ("Prepared By (Analyst):", f"{signoff_data.get('preparer_name')} [ID: {signoff_data.get('preparer_id')}]", f"Signed: {signoff_data.get('preparer_date')} ({sig_status})"),
            ("Reviewed By (Controller / Manager):", f"{signoff_data.get('reviewer_name')} [{signoff_data.get('reviewer_title')}]", f"Reviewed: {signoff_data.get('reviewer_date')} | Verdict: {signoff_data.get('review_verdict')}"),
            ("Audit Workpaper Reference:", signoff_data.get("wp_ref"), f"Status: {signoff_data.get('review_verdict')}"),
            ("Supervisory Audit Notes:", signoff_data.get("audit_comments", "None noted"), "Document Status: Final Closed")
        ]

        for c1, c2, c3 in signoff_entries:
            ws_summary.append([c1, c2, c3])

        for r in range(sign_off_header_row + 1, sign_off_header_row + 5):
            for c in range(1, 4):
                cell = ws_summary.cell(row=r, column=c)
                cell.font = bold_font if c == 1 else regular_font
                cell.border = thin_border

        ws_summary.column_dimensions["A"].width = 64
        ws_summary.column_dimensions["B"].width = 60
        ws_summary.column_dimensions["C"].width = 65

        data_sheets = {
            "Deterministic Matches": m_copy,
            "Fuzzy Candidates": f_copy,
            "Exceptions Source A": a_copy,
            "Exceptions Source B": b_copy
        }

        for sheet_title, df_content in data_sheets.items():
            df_content.to_excel(writer, sheet_name=sheet_title, index=False)
            ws = writer.sheets[sheet_title]
            ws.views.sheetView[0].showGridLines = True
            ws.freeze_panes = "A2"

            for col_idx in range(1, len(df_content.columns) + 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.fill = navy_header_fill
                cell.font = white_bold_font
                cell.alignment = Alignment(horizontal="center", vertical="center")

            for row in ws.iter_rows(min_row=2, max_row=len(df_content) + 1, min_col=1, max_col=len(df_content.columns)):
                for cell in row:
                    cell.border = thin_border
                    cell.font = regular_font
                    if isinstance(cell.value, (int, float)):
                        cell.number_format = "#,##0.00"

            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    val_str = str(cell.value or "")
                    max_len = max(max_len, len(val_str))
                ws.column_dimensions[col_letter].width = min(max(max_len + 4, 15), 55)

    return output.getvalue()

# ==============================================================================
# --- SIDEBAR CONTROLS ---
# ==============================================================================
st.sidebar.markdown("### ⚡ AutoReconFlow")
st.sidebar.markdown("**Enterprise B2B Edition v2.3**")
st.sidebar.caption("SOX 404 / ISA-500, Indian GST & International VAT Architecture")
st.sidebar.divider()

uploaded_logo = st.sidebar.file_uploader("Upload Corporate Logo", type=["png", "jpg", "jpeg"], key="company_logo")
company_name = st.sidebar.text_input("Entity / CA Firm Legal Name", value="Global Enterprise Solutions Ltd.", key="cfg_comp_name")
company_address = st.sidebar.text_area("Registered Office Address", value="Tower 4, Financial District, Gurugram, India - 122002", key="cfg_comp_addr", height=65)
company_uid = st.sidebar.text_input("Tax / Entity ID (GSTIN / VAT / CIN)", value="CIN: U72200HR2026PTC109823", key="cfg_comp_uid")
st.sidebar.divider()

# Zero-Token SMTP Mailbox Settings with Live Ping
with st.sidebar.expander("📬 Automated Email Alert Settings", expanded=False):
    st.caption("Zero-token SMTP engine. Compatible with Gmail App Passwords, Outlook, and Brevo.")
    enable_smtp = st.checkbox("Enable Live Email Dispatch", value=False, key="smtp_enabled")
    
    default_host = st.secrets.get("SMTP_HOST", "smtp.gmail.com") if hasattr(st, "secrets") else "smtp.gmail.com"
    default_port = str(st.secrets.get("SMTP_PORT", "465")) if hasattr(st, "secrets") else "465"
    default_user = st.secrets.get("SMTP_USER", "") if hasattr(st, "secrets") else ""
    default_pass = st.secrets.get("SMTP_PASSWORD", "") if hasattr(st, "secrets") else ""

    smtp_host = st.text_input("SMTP Server Host", value=default_host, key="smtp_host")
    smtp_port = st.text_input("SMTP Port (465=SSL / 587=TLS)", value=default_port, key="smtp_port")
    smtp_user = st.text_input("Sender Email Address", value=default_user, placeholder="yourname@gmail.com", key="smtp_user")
    smtp_pass = st.text_input("Sender App Password", value=default_pass, type="password", placeholder="16-character app password", key="smtp_pass")

    smtp_cfg = {
        "enabled": enable_smtp,
        "host": smtp_host,
        "port": smtp_port,
        "user": smtp_user,
        "password": smtp_pass
    }

    test_target = st.text_input("Test Target Email", value=smtp_user or "", placeholder="recipient@domain.com", key="smtp_test_target")
    if st.button("🧪 Send Test Ping Email", use_container_width=True):
        if not enable_smtp:
            st.warning("Please check 'Enable Live Email Dispatch' first.")
        else:
            with st.spinner("Connecting to SMTP server..."):
                ok, feedback = send_workflow_email(
                    smtp_cfg, 
                    test_target, 
                    "AutoReconFlow - SMTP Connection Verified", 
                    "This confirms your AutoReconFlow email trigger configuration is working properly.",
                    build_html_email("SMTP Connection Verified", "Your AutoReconFlow email alerts are operational.", {"Status": "Operational", "Tested At": datetime.datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S UTC")})
                )
                if ok:
                    st.success("✅ Test email dispatched successfully! Check your inbox.")
                else:
                    st.error(f"❌ {feedback}")

st.sidebar.divider()

# In-App License Verification Widget
with st.sidebar.expander("🔑 License Activation & Verification", expanded=True):
    user_email_input = st.text_input("Registered Email", placeholder="buyer@clientdomain.com", key="lic_user_email")
    user_key_input = st.text_input("License Key", placeholder="ARF-PRAC-XXXX-XXXX-XXXX", key="lic_user_key")
    if user_email_input and user_key_input:
        is_valid, active_tier = verify_license_key(user_email_input, user_key_input)
        if is_valid:
            st.success(f"✅ Verified: {active_tier}")
            st.session_state["licensed"] = True
            st.session_state["license_tier"] = active_tier
        else:
            st.error("❌ Invalid License Key or Email Mismatch.")
            st.session_state["licensed"] = False
    else:
        st.caption("Enter credentials to unlock full audit deliverables.")

st.sidebar.divider()
st.sidebar.markdown("### 💼 Commercial Licensing")
st.sidebar.markdown(
    """
    * **CA / Tax Firm License (₹7,499 / $99/Year)**: Unlimited GSTINs, GSTR-2B matching, and 4-bucket Excel exports.
    * **Business & Tax Suite (₹16,999 / $249/Year)**: Cross-border VAT RCM audit, subledger-to-bank engine, and full source code access.
    * **Enterprise Turnkey (₹39,999 / $499/Year)**: Private server deployment, multi-user database integration, and compliance onboarding.
    """
)
st.sidebar.link_button("💳 Purchase Commercial License", "https://rzp.io/rzp/G2XPtSpy", use_container_width=True, type="primary")
st.sidebar.caption("Pure Python • 0 Token Costs • 100% On-Premise Data Privacy")

# ==============================================================================
# --- MAIN APPLICATION WORKSPACE & STATUTORY NAVIGATION TABS ---
# ==============================================================================
st.markdown("""
<div class="brand-hero">
    <h1>⚖️ AutoReconFlow — Enterprise Financial Reconciliation, Tax & Audit Engine</h1>
    <p>Statutory-compliant reconciliation engine covering Indian GST Section 16(2)(aa), International VAT RCM (HMRC MTD / EU Directive), and SOX 404 / ISA-500 Four-Eyes Governance.</p>
</div>
""", unsafe_allow_html=True)

tab_gst, tab_vat, tab_universal = st.tabs([
    "🇮🇳 Indian GST ITC Reconciler [Sec 16(2)(aa) & Rule 36(4)]",
    "🌍 Cross-Border VAT & Reverse Charge (RCM) [HMRC Sec 8 / EU Directive Art 196]",
    "🌐 Universal Bank vs. General Ledger [SOX 404 / PCAOB AS 2201 / ISA-500]"
])

# ==============================================================================
# --- TAB 1: INDIAN GST ITC (PURCHASE REGISTER VS. GSTR-2B) ---
# ==============================================================================
with tab_gst:
    st.markdown("""
    <div class="compliance-badge-card">
        <strong>🏛️ Statutory Compliance Framework (India GST):</strong><br>
        • <strong>CGST Act 2017 Section 16(2)(aa)</strong>: Legal prohibition against claiming Input Tax Credit unless the invoice has been furnished by the supplier in GSTR-1/IFF and reflected in the recipient's auto-drafted GSTR-2B.<br>
        • <strong>CGST Rule 36(4)</strong>: Mandatory reconciliation before filing monthly GSTR-3B.<br>
        • <strong>CGST Section 50 Scrutiny Risk</strong>: Claiming unreflected ITC attracts automated GST scrutiny notices with 18% per annum interest liability.
    </div>
    """, unsafe_allow_html=True)

    SAMPLE_PR = """Supplier_GSTIN,Supplier_Name,Invoice_Number,Invoice_Date,Taxable_Value,IGST,CGST,SGST
07AAAAA1111A1Z1,Apex Technologies Ltd,INV/2026/0101,2026-10-01,100000.00,18000.00,0.00,0.00
06BBBBB2222B2Z2,BlueSky Logistics,BSL-4580,2026-10-02,50000.00,0.00,4500.00,4500.00
27CCCCC3333C3Z3,CloudCore Infotech,CC-991,2026-10-03,80000.00,14400.00,0.00,0.00
08DDDDD4444D4Z4,Delta Industrial Supplies,INV-0044,2026-10-04,30000.00,0.00,2700.00,2700.00
07EEEEE5555E5Z5,Echo Communications,EC/26/55,2026-10-05,45000.00,8100.00,0.00,0.00
29FFFFF6666F6Z6,Falcon Freight Systems,FF-880,2026-10-06,60000.00,10800.00,0.00,0.00"""

    SAMPLE_2B = """GSTIN_of_Supplier,Trade_Name,Invoice_number,Invoice_date,Taxable_Value_Rs,Integrated_Tax_Rs,Central_Tax_Rs,State_UT_Tax_Rs
07AAAAA1111A1Z1,Apex Technologies Ltd,INV-2026-101,01-10-2026,100000.00,18000.00,0.00,0.00
06BBBBB2222B2Z2,BlueSky Logistics Pvt Ltd,BSL4580,02-10-2026,50000.00,0.00,4500.00,4500.00
27CCCCC3333C3Z3,CloudCore Infotech,CC991,03-10-2026,80000.00,14400.00,0.00,0.00
08DDDDD4444D4Z4,Delta Industrial,44,04-10-2026,30000.00,0.00,2700.00,2700.00
07EEEEE5555E5Z5,Echo Communications,EC-26-55,05-10-2026,45000.00,7500.00,0.00,0.00
03GGGGG7777G7Z7,Global Spares Inc,GS-1004,06-10-2026,25000.00,4500.00,0.00,0.00"""

    use_gst_mock = st.checkbox("🧪 Load Real-World GST Test Files (With invoice format variations)", value=True, key="chk_gst_mock")
    df_pr = pd.read_csv(io.StringIO(SAMPLE_PR)) if use_gst_mock else None
    df_2b = pd.read_csv(io.StringIO(SAMPLE_2B)) if use_gst_mock else None

    if not use_gst_mock:
        u1, u2 = st.columns(2)
        with u1:
            fpr = st.file_uploader("Upload Purchase Register (Books / Tally)", type=["csv", "xlsx"], key="up_pr")
            if fpr: df_pr = robust_read_file(fpr)
        with u2:
            f2b = st.file_uploader("Upload GSTR-2B Statement (GST Portal)", type=["csv", "xlsx"], key="up_2b")
            if f2b: df_2b = robust_read_file(f2b)

    if df_pr is not None and df_2b is not None:
        st.success(f"✓ Loaded Purchase Register ({len(df_pr)} rows) and GSTR-2B ({len(df_2b)} rows)")
        
        with st.expander("⚙️ Field Mapping & Statutory Tolerance", expanded=True):
            mc1, mc2 = st.columns(2)
            cols_pr = list(df_pr.columns)
            cols_2b = list(df_2b.columns)

            def get_idx(cols, patterns, fallback=0):
                for idx, c in enumerate(cols):
                    clean_c = str(c).lower().replace(" ", "_")
                    if any(p in clean_c for p in patterns):
                        return idx
                return min(fallback, len(cols) - 1)

            with mc1:
                st.markdown("**Purchase Register Columns**")
                pr_gstin = st.selectbox("Supplier GSTIN", cols_pr, index=get_idx(cols_pr, ["gstin", "supplier_gst", "tin"]), key="pr_g")
                pr_inv = st.selectbox("Invoice Number", cols_pr, index=get_idx(cols_pr, ["inv", "bill", "doc"]), key="pr_i")
                pr_taxable = st.selectbox("Taxable Value", cols_pr, index=get_idx(cols_pr, ["taxable", "basic", "amount"]), key="pr_tx")
                pr_igst = st.selectbox("IGST Amount", cols_pr, index=get_idx(cols_pr, ["igst", "integrated"]), key="pr_ig")
                pr_cgst = st.selectbox("CGST Amount", cols_pr, index=get_idx(cols_pr, ["cgst", "central"]), key="pr_cg")
                pr_sgst = st.selectbox("SGST Amount", cols_pr, index=get_idx(cols_pr, ["sgst", "state"]), key="pr_sg")

            with mc2:
                st.markdown("**GSTR-2B Columns**")
                g2b_gstin = st.selectbox("Supplier GSTIN", cols_2b, index=get_idx(cols_2b, ["gstin", "supplier_gst"]), key="g2b_g")
                g2b_inv = st.selectbox("Invoice Number", cols_2b, index=get_idx(cols_2b, ["inv", "doc"]), key="g2b_i")
                g2b_taxable = st.selectbox("Taxable Value", cols_2b, index=get_idx(cols_2b, ["taxable", "value"]), key="g2b_tx")
                g2b_igst = st.selectbox("IGST Amount", cols_2b, index=get_idx(cols_2b, ["igst", "integrated"]), key="g2b_ig")
                g2b_cgst = st.selectbox("CGST Amount", cols_2b, index=get_idx(cols_2b, ["cgst", "central"]), key="g2b_cg")
                g2b_sgst = st.selectbox("SGST Amount", cols_2b, index=get_idx(cols_2b, ["sgst", "state"]), key="g2b_sg")

            gst_tol = st.slider("Tax Rounding Tolerance (₹)", 0.00, 10.00, 2.00, step=0.50, key="gst_tol_slider")

        if st.button("🚀 Run GSTR-2B vs. Purchase Register Matching", type="primary", use_container_width=True, key="btn_run_gst"):
            map_pr = {'gstin': pr_gstin, 'inv': pr_inv, 'taxable': pr_taxable, 'igst': pr_igst, 'cgst': pr_cgst, 'sgst': pr_sgst}
            map_2b = {'gstin': g2b_gstin, 'inv': g2b_inv, 'taxable': g2b_taxable, 'igst': g2b_igst, 'cgst': g2b_cgst, 'sgst': g2b_sgst}
            b1, b2, b3, b4 = execute_gst_itc_reconciliation(df_pr, df_2b, map_pr, map_2b, tolerance=gst_tol)
            st.session_state["gst_res"] = (b1, b2, b3, b4)

        if "gst_res" in st.session_state:
            b1, b2, b3, b4 = st.session_state["gst_res"]
            st.markdown("#### 📊 GSTR-3B Statutory ITC Audit Buckets")
            k1, k2, k3, k4 = st.columns(4)
            k1.metric("Bucket 1: Direct Matches", f"{len(b1)} invoices", delta="Eligible to Claim")
            k2.metric("Bucket 2: Missing in Books", f"{len(b2)} invoices", delta="Verify & Book")
            k3.metric("Bucket 3: Missing in 2B", f"{len(b3)} invoices", delta="-ITC Blocked", delta_color="inverse")
            k4.metric("Bucket 4: Tax Mismatches", f"{len(b4)} invoices", delta="Requires Review", delta_color="inverse")

            wb_bytes = generate_gst_itc_excel_workpaper(b1, b2, b3, b4, company_name, company_uid)
            st.download_button(
                label="📥 Download Master GSTR-3B Audit Workpaper (.xlsx)",
                data=wb_bytes,
                file_name=f"GSTR2B_Reconciliation_Workpaper_{datetime.datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                type="primary",
                use_container_width=True,
                key="btn_dl_gst_wb"
            )

            t1, t2, t3, t4 = st.tabs([
                "✅ Bucket 1: Eligible ITC (Matched)",
                "⚠️ Bucket 2: Available in 2B (Missing in Books)",
                "🚫 Bucket 3: Missing in 2B (Blocked ITC)",
                "🔍 Bucket 4: Value & Tax Discrepancies"
            ])
            with t1: st.dataframe(b1, use_container_width=True)
            with t2: st.dataframe(b2, use_container_width=True)
            with t3: 
                st.dataframe(b3, use_container_width=True)
                st.warning("⚠️ Suppliers listed above have not filed GSTR-1. Claiming this ITC risks automated scrutiny notices under Section 16(2)(aa).")
                
                if len(b3) > 0:
                    st.markdown("##### ✉️ Automated Vendor Notice Trigger Desk")
                    unique_suppliers = b3['Supplier_Name'].dropna().unique().tolist() if 'Supplier_Name' in b3.columns else []
                    if unique_suppliers:
                        sel_supp = st.selectbox("Select Non-Filing Supplier to Issue Notice:", unique_suppliers, key="sel_supp_gst")
                        supp_missing = b3[b3['Supplier_Name'] == sel_supp]
                        col_e1, col_e2 = st.columns([2, 1])
                        with col_e1:
                            target_supp_email = st.text_input("Supplier Accounts Email", placeholder="accounts@supplierdomain.com", key="supp_email_input")
                        notice_html, blocked_sum = build_vendor_gst_notice(company_name, company_uid, sel_supp, supp_missing)
                        with st.expander(f"👁️ Preview Statutory Notice Draft (Blocked Tax: ₹{blocked_sum:,.2f})"):
                            st.components.v1.html(notice_html, height=340, scrolling=True)
                        if st.button("📤 Send Statutory GST Non-Filing Notice", type="primary", key="btn_send_supp_gst"):
                            if not target_supp_email or "@" not in target_supp_email:
                                st.error("Please provide a valid supplier email.")
                            else:
                                with st.spinner("Dispatching via SMTP..."):
                                    ok, info = send_workflow_email(
                                        smtp_cfg, target_supp_email, 
                                        f"URGENT: GSTR-1 Non-Filing Notice — {company_name}", 
                                        f"Blocked ITC: ₹{blocked_sum}", notice_html
                                    )
                                    if ok: st.success("✓ Notice dispatched successfully!")
                                    else: st.error(f"Failed: {info}")
            with t4: st.dataframe(b4, use_container_width=True)

# ==============================================================================
# --- TAB 2: INTERNATIONAL TWO-SIDED VAT & REVERSE CHARGE (RCM) ---
# ==============================================================================
with tab_vat:
    st.markdown("""
    <div class="compliance-badge-card">
        <strong>🏛️ Statutory Compliance Framework (International Cross-Border VAT):</strong><br>
        • <strong>Two-Sided Reconciliation</strong>: Matches Internal Purchase Books (Source A) against Overseas Supplier Invoices / Statements (Source B).<br>
        • <strong>UK VAT Act 1994, Section 8 & HMRC MTD</strong>: Mandates self-assessment of output tax on B2B cross-border services while allowing qualifying input tax deduction.<br>
        • <strong>EU VAT Directive 2006/112/EC (Article 196)</strong>: Reverse Charge mechanism on cross-border transactions (Net Tax Payable = 0.00).<br>
        • <strong>Statutory Risk Flag</strong>: Flags instances where foreign suppliers erroneously charge foreign VAT on B2B reverse-charge transactions.
    </div>
    """, unsafe_allow_html=True)

    SAMPLE_BOOKS_VAT = """Invoice_Ref,Supplier_Name,Country,Supplier_VAT_ID,Net_Amount_GBP,Is_Reverse_Charge,Self_Assessed_VAT_GBP
INV-UK-901,London Cloud Hosting,UK,GB123456789,12000.00,NO,0.00
INV-US-882,San Francisco SaaS Corp,USA,US-DEL-9921,45000.00,YES,9000.00
INV-EU-741,Munich Enterprise Software,DE,DE987654321,18000.00,YES,900.00
INV-SG-109,Singapore AI Solutions,SG,SG-UEN-2026,22000.00,YES,4400.00
INV-US-550,Boston Tech Consulting,USA,US-MA-4412,8500.00,YES,1700.00"""

    SAMPLE_EXT_VAT = """Invoice_Number,Supplier_Legal_Name,Vendor_Tax_ID,Billed_Net_GBP,Charged_VAT_GBP,Billing_Currency
INV901,London Cloud Hosting Ltd,GB123456789,12000.00,2400.00,GBP
INV-US-882,San Francisco SaaS Corp,US-DEL-9921,45000.00,0.00,GBP
INV-EU-741,Munich Enterprise Software,DE987654321,18000.00,0.00,GBP
INV-SG-109,Singapore AI Solutions,SG-UEN-2026,22000.00,4400.00,GBP
INV-FR-301,Paris Cloud Analytics,FR88991122,14000.00,0.00,EUR"""

    use_intl_mock = st.checkbox("🧪 Load Real-World Two-Sided VAT Test Datasets", value=True, key="chk_intl_mock")
    df_books_vat = pd.read_csv(io.StringIO(SAMPLE_BOOKS_VAT)) if use_intl_mock else None
    df_ext_vat = pd.read_csv(io.StringIO(SAMPLE_EXT_VAT)) if use_intl_mock else None

    if not use_intl_mock:
        v_col1, v_col2 = st.columns(2)
        with v_col1:
            st.subheader("Source A: Internal Purchase Books / ERP")
            f_bks = st.file_uploader("Upload Internal Books (CSV / Excel)", type=["csv", "xlsx"], key="up_vat_books")
            if f_bks: df_books_vat = robust_read_file(f_bks)
        with v_col2:
            st.subheader("Source B: Supplier Statement / Tax Extract")
            f_ext = st.file_uploader("Upload Supplier / External Statement (CSV / Excel)", type=["csv", "xlsx"], key="up_vat_ext")
            if f_ext: df_ext_vat = robust_read_file(f_ext)

    if df_books_vat is not None and df_ext_vat is not None:
        st.success(f"✓ Source A Books ({len(df_books_vat)} rows) and Source B External ({len(df_ext_vat)} rows) loaded.")

        with st.expander("⚙️ VAT Configuration & Two-Sided Column Mapping", expanded=True):
            r1, r2, r3 = st.columns(3)
            with r1:
                statutory_vat = st.number_input("Domestic Statutory VAT Rate (%)", value=20.0, step=0.5, key="stat_vat_rate")
            with r2:
                audit_currency = st.selectbox("Reporting Base Currency", ["GBP (£)", "EUR (€)", "USD ($)", "AED (د.إ)", "SAR (﷼)"], key="vat_curr")
            with r3:
                rcm_tolerance = st.number_input("Allowed Variance Threshold", value=1.0, step=0.1, key="vat_tol")

            cols_bks = list(df_books_vat.columns)
            cols_ext = list(df_ext_vat.columns)

            m_b1, m_b2 = st.columns(2)
            with m_b1:
                st.markdown("**Source A (Internal Books)**")
                map_bks_inv = st.selectbox("Invoice Ref Column", cols_bks, index=0, key="mb_inv")
                map_bks_taxid = st.selectbox("Supplier VAT/Tax ID Column", cols_bks, index=min(3, len(cols_bks)-1), key="mb_tid")
                map_bks_net = st.selectbox("Net Taxable Amount Column", cols_bks, index=min(4, len(cols_bks)-1), key="mb_net")
                map_bks_rcm = st.selectbox("Is Reverse Charge Flag Column", cols_bks, index=min(5, len(cols_bks)-1), key="mb_rcm")
                map_bks_self = st.selectbox("Self-Assessed Deemed VAT Column", cols_bks, index=min(6, len(cols_bks)-1), key="mb_slf")
            with m_b2:
                st.markdown("**Source B (External Supplier Statement)**")
                map_ext_inv = st.selectbox("Invoice Number Column", cols_ext, index=0, key="me_inv")
                map_ext_taxid = st.selectbox("Vendor VAT/Tax ID Column", cols_ext, index=min(2, len(cols_ext)-1), key="me_tid")
                map_ext_net = st.selectbox("Billed Net Amount Column", cols_ext, index=min(3, len(cols_ext)-1), key="me_net")
                map_ext_chg = st.selectbox("Charged VAT Column", cols_ext, index=min(4, len(cols_ext)-1), key="me_chg")

        if st.button("🚀 Run Two-Sided Cross-Border VAT & RCM Reconciliation", type="primary", use_container_width=True, key="btn_run_two_sided_vat"):
            mb = {'invoice': map_bks_inv, 'vat_id': map_bks_taxid, 'net_amt': map_bks_net, 'is_rcm': map_bks_rcm, 'self_assessed': map_bks_self}
            me = {'invoice': map_ext_inv, 'vat_id': map_ext_taxid, 'net_amt': map_ext_net, 'charged_vat': map_ext_chg}
            
            b1_v, b2_v, b3_v, b4_v = execute_two_sided_vat_rcm_reconciliation(df_books_vat, df_ext_vat, mb, me, applicable_vat_rate=statutory_vat, tolerance=rcm_tolerance)
            st.session_state["two_sided_vat_res"] = (b1_v, b2_v, b3_v, b4_v)

        if "two_sided_vat_res" in st.session_state:
            b1_v, b2_v, b3_v, b4_v = st.session_state["two_sided_vat_res"]

            st.markdown("#### 📊 Cross-Border VAT Governance Assertion Metrics")
            k1, k2, k3, k4 = st.columns(4)
            k1.metric("Bucket 1: Direct Matches", f"{len(b1_v)} lines", delta="RCM Audit Cleared")
            k2.metric("Bucket 2: Missing in Books", f"{len(b2_v)} lines", delta="Unrecorded Liability", delta_color="inverse")
            k3.metric("Bucket 3: Missing in Statement", f"{len(b3_v)} lines", delta="Accrued in Books", delta_color="inverse")
            k4.metric("Bucket 4: Rate & Tax Mismatches", f"{len(b4_v)} lines", delta="Foreign VAT / Discrepancy", delta_color="inverse")

            # Excel Export for 4-bucket VAT Workpaper
            output_vat = io.BytesIO()
            with pd.ExcelWriter(output_vat, engine="openpyxl") as vat_writer:
                b1_v.to_excel(vat_writer, sheet_name="Bucket 1 - Cleared RCM", index=False)
                b2_v.to_excel(vat_writer, sheet_name="Bucket 2 - Missing in Books", index=False)
                b3_v.to_excel(vat_writer, sheet_name="Bucket 3 - Missing in Ext", index=False)
                b4_v.to_excel(vat_writer, sheet_name="Bucket 4 - Discrepancies", index=False)
            
            st.download_button(
                label="📥 Download Two-Sided VAT & RCM Audit Workpaper (.xlsx)",
                data=output_vat.getvalue(),
                file_name=f"Cross_Border_VAT_Two_Sided_Workpaper_{datetime.datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                type="primary",
                use_container_width=True,
                key="btn_dl_two_sided_vat"
            )

            vt1, vt2, vt3, vt4 = st.tabs([
                "✅ Bucket 1: Cleared RCM Matches",
                "⚠️ Bucket 2: Missing in Books (Unrecorded)",
                "🔍 Bucket 3: Missing in External Statement",
                "🚫 Bucket 4: Rate & Tax Discrepancies (Unlawful Foreign Tax)"
            ])
            with vt1: st.dataframe(b1_v, use_container_width=True)
            with vt2: st.dataframe(b2_v, use_container_width=True)
            with vt3: st.dataframe(b3_v, use_container_width=True)
            with vt4: 
                st.dataframe(b4_v, use_container_width=True)
                st.warning("⚠️ Any overseas vendor listed in Bucket 4 who charged domestic VAT on an RCM service must be issued an amendment request for a zero-rated invoice.")

# ==============================================================================
# --- TAB 3: UNIVERSAL BANK VS. GENERAL LEDGER (COMPLETE RESTORATION) ---
# ==============================================================================
with tab_universal:
    st.markdown("""
    <div class="compliance-badge-card">
        <strong>🏛️ Statutory Compliance Framework (Internal Governance):</strong><br>
        • <strong>Sarbanes-Oxley (SOX) Act Section 404 & PCAOB AS 2201</strong>: Segregation of Duties (SoD) requiring documented Preparer verification and an independent Supervisory Reviewer approval loop.<br>
        • <strong>ISA-500 (International Standard on Auditing)</strong>: Audit trail integrity, non-repudiation, line-item query threads, and cryptographic SHA-256 digital immutability hashing.
    </div>
    """, unsafe_allow_html=True)

    with st.expander("📖 Comprehensive Procedure & Configuration Manual (SOX 404 / ISA-500 Standard)", expanded=False):
        st.markdown("""
        ### 🏛️ Section 1: Corporate Entity & Workpaper Branding Setup
        AutoReconFlow stamps corporate identity metadata across all audit deliverables (Multi-tab Excel workpapers and signed PDF certificates).
        
        ### 📬 Section 2: Automated Email Dispatch & SMTP Mapping
        Zero-token built-in SMTP engine (`smtplib`). No third-party paid API tokens required.
        
        ### 🔄 Section 3: Four-Eyes Governance & Line-Item Audit Query SOP
        Enforces strict segregation of duties (SoD) under PCAOB AS 2201 and SOX 404.
        """)

    SAMPLE_A_CSV = """Invoice_ID,Vendor_Name,Amount,Posting_Date
INV-2026-101,Apex Tech Solutions,12500.00,2026-10-01
INV-2026-102,BlueSky Logistics,8450.50,2026-10-02
INV-2026-103,CloudCore Systems,19200.00,2026-10-03
INV-2026-104,Delta Industrial Supplies,4300.00,2026-10-04
INV-2026-105,Echo Media Group,15750.25,2026-10-05
INV-2026-106,Falcon Freight Services,6200.00,2026-10-06
INV-2026-107,Global Consulting Partners,31000.00,2026-10-07
INV-2026-108,Horizon Office Supplies,1450.00,2026-10-08
INV-2026-109,Infinite Cloud Storage,9800.00,2026-10-09
INV-2026-110,Jupiter Legal Advisory,22500.00,2026-10-10"""

    SAMPLE_B_CSV = """Transaction_Ref,Description,Settlement_Amount,Value_Date
INV2026101,Apex Tech Solutions LLC,12500.00,2026-10-02
INV2026102,BlueSky Logistics,8450.50,2026-10-03
INV-2026-103,CloudCore Sys Inc,19200.00,2026-10-04
INV-2026-104-MISC,Delta Industrial Supplies,4300.00,2026-10-05
INV-2026-105,Echo Media,15750.25,2026-10-06
INV-2026-999,Vanguard Security,5100.00,2026-10-06
INV-2026-107,Global Consulting,31000.00,2026-10-08
INV-2026-108,Horizon Office Supplies,1449.50,2026-10-09
INV-2026-888,QuickFleet Courier,3200.00,2026-10-10
INV-2026-110,Jupiter Legal,22500.00,2026-10-11"""

    use_mock_univ = st.checkbox("Load Built-in Sample Datasets (Source A & Source B)", value=True, key="chk_mock_univ")
    df_a = pd.read_csv(io.StringIO(SAMPLE_A_CSV)) if use_mock_univ else None
    df_b = pd.read_csv(io.StringIO(SAMPLE_B_CSV)) if use_mock_univ else None

    if not use_mock_univ:
        col1, col2 = st.columns(2)
        with col1:
            st.subheader("Source A (Ledger / Internal Records)")
            file_a = st.file_uploader("Upload CSV, XLSX, or Bank Statement PDF", type=["csv", "xlsx", "xls", "pdf"], key="file_a")
            if file_a: df_a = robust_read_file(file_a)
        with col2:
            st.subheader("Source B (Bank / Tax Portal / External Records)")
            file_b = st.file_uploader("Upload CSV, XLSX, or Bank Statement PDF", type=["csv", "xlsx", "xls", "pdf"], key="file_b")
            if file_b: df_b = robust_read_file(file_b)

    def get_default_index(columns, keywords, fallback_index=0):
        for idx, col in enumerate(columns):
            clean_col = str(col).lower().replace(" ", "_")
            if any(kw in clean_col for kw in keywords):
                return idx
        return min(fallback_index, len(columns) - 1)

    if df_a is not None and df_b is not None:
        cols_a = list(df_a.columns)
        cols_b = list(df_b.columns)
        st.success(f"Loaded Source A ({len(df_a):,} rows, {len(cols_a)} cols) and Source B ({len(df_b):,} rows, {len(cols_b)} cols)")

        with st.expander("🛠️ Configure Column Mappings", expanded=True):
            m1, m2 = st.columns(2)
            with m1:
                st.markdown("**Source A (Internal Books)**")
                idx_id_a = get_default_index(cols_a, ["invoice", "inv", "doc", "ref", "id"], 0)
                idx_party_a = get_default_index(cols_a, ["vendor", "party", "name", "desc"], min(1, len(cols_a)-1))
                idx_amt_a = get_default_index(cols_a, ["amount", "total", "val", "debit"], min(2, len(cols_a)-1))
                id_a = st.selectbox("Document / Txn ID Column", cols_a, index=idx_id_a, key="map_id_a")
                party_a = st.selectbox("Party / Narrative Column", cols_a, index=idx_party_a, key="map_party_a")
                amt_a = st.selectbox("Amount Column", cols_a, index=idx_amt_a, key="map_amt_a")
            with m2:
                st.markdown("**Source B (Bank / Portal)**")
                idx_id_b = get_default_index(cols_b, ["transaction", "ref", "invoice", "inv", "id"], 0)
                idx_party_b = get_default_index(cols_b, ["description", "narrative", "party", "particulars"], min(1, len(cols_b)-1))
                idx_amt_b = get_default_index(cols_b, ["settlement", "amount", "total", "net"], min(2, len(cols_b)-1))
                id_b = st.selectbox("Document / Txn ID Column", cols_b, index=idx_id_b, key="map_id_b")
                party_b = st.selectbox("Party / Narrative Column", cols_b, index=idx_party_b, key="map_party_b")
                amt_b = st.selectbox("Amount Column", cols_b, index=idx_amt_b, key="map_amt_b")

        tolerance = st.slider("Allowed Amount Tolerance (Currency Units)", 0.00, 5.00, 1.00, step=0.10, key="univ_tol_slider")

        if st.button("🚀 Run Automated Reconciliation", type="primary", key="btn_run_univ"):
            map_a = {"id": id_a, "party": party_a, "amt": amt_a}
            map_b = {"id": id_b, "party": party_b, "amt": amt_b}
            pl_a = pl.from_pandas(df_a)
            pl_b = pl.from_pandas(df_b)

            with st.spinner("Processing through 3-tier engine..."):
                try:
                    matched, fuzzy, un_a, un_b = execute_high_volume_reconciliation(
                        pl_a, pl_b, map_a, map_b, tolerance=tolerance
                    )
                except Exception as e:
                    st.error(f"Reconciliation error: {str(e)}")
                    st.stop()

            st.session_state["recon_results"] = {
                "matched_pd": matched.to_pandas(),
                "fuzzy_pd": fuzzy.to_pandas(),
                "un_a_pd": un_a.to_pandas(),
                "un_b_pd": un_b.to_pandas(),
                "matched_pl": matched,
                "fuzzy_pl": fuzzy,
                "un_a_pl": un_a,
                "un_b_pl": un_b,
                "tolerance": tolerance
            }
            st.session_state["preparer_completed"] = False
            st.session_state["reviewer_approved"] = False
            st.session_state["audit_query_thread"] = {
                "status": "NONE",
                "item_id": "",
                "reviewer_query": "",
                "reviewer_query_time": "",
                "preparer_reply": "",
                "preparer_reply_time": ""
            }

        if "recon_results" in st.session_state:
            res = st.session_state["recon_results"]
            matched_pd = res["matched_pd"]
            fuzzy_pd = res["fuzzy_pd"]
            un_a_pd = res["un_a_pd"]
            un_b_pd = res["un_b_pd"]

            st.markdown("#### 📊 Internal Control Summary & Transaction Metrics")
            k1, k2, k3, k4 = st.columns(4)
            k1.metric("Tier-1 Exact Matches", f"{len(matched_pd):,} records")
            k2.metric("Tier-2 Variance Review", f"{len(fuzzy_pd):,} records")
            k3.metric("Source A Ledger Exceptions", f"{len(un_a_pd):,} records")
            k4.metric("Source B External Exceptions", f"{len(un_b_pd):,} records")

            is_prep_done = st.session_state.get("preparer_completed", False)
            is_rev_approved = st.session_state.get("reviewer_approved", False)
            if "audit_query_thread" not in st.session_state:
                st.session_state["audit_query_thread"] = {
                    "status": "NONE",
                    "item_id": "",
                    "reviewer_query": "",
                    "reviewer_query_time": "",
                    "preparer_reply": "",
                    "preparer_reply_time": ""
                }
            query_state = st.session_state["audit_query_thread"]

            # --- FOUR-EYES GOVERNANCE SIGN-OFF WORKFLOW ---
            with st.expander("✍️ Four-Eyes Internal Governance Sign-Off & Approval Workflow", expanded=True):
                so_col1, so_col2 = st.columns(2)
                now_str = datetime.datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S UTC")

                # STAGE 1: PREPARER PANEL
                with so_col1:
                    st.markdown("##### 👤 Stage 1: Preparer (Analyst / Senior Accountant)")
                    prep_name = st.text_input("Preparer Full Name", value="Senior Recon Analyst", key="so_prep_name", disabled=is_prep_done)
                    prep_id = st.text_input("Employee / System ID", value="FIN-ACC-8921", key="so_prep_id", disabled=is_prep_done)
                    prep_email = st.text_input("Preparer Email Address", value="analyst@company.com", key="so_prep_email", disabled=is_prep_done)
                    rev_target_email = st.text_input("Reviewer Notification Email", value="controller@company.com", key="so_rev_target_email", disabled=is_prep_done)
                    prep_date = st.text_input("Preparation Date & Time (UTC)", value=st.session_state.get("saved_prep_date", now_str), key="so_prep_date", disabled=is_prep_done)
                    prep_signed = st.checkbox("I verify that ledger entries, tolerances, and files are audited.", value=True, key="so_prep_signed", disabled=is_prep_done)

                    if not is_prep_done:
                        if st.button("✅ Mark Preparer Task as Complete & Notify Reviewer", type="primary", use_container_width=True, key="btn_prep_submit"):
                            if not prep_signed:
                                st.error("Preparer verification box must be checked before submission.")
                            else:
                                st.session_state["preparer_completed"] = True
                                st.session_state["saved_prep_name"] = prep_name
                                st.session_state["saved_prep_id"] = prep_id
                                st.session_state["saved_prep_email"] = prep_email
                                st.session_state["saved_rev_target_email"] = rev_target_email
                                st.session_state["saved_prep_date"] = prep_date
                                st.session_state["saved_prep_signed"] = prep_signed

                                plain_body = f"Reconciliation workpaper submitted by {prep_name} ({prep_id}) at {prep_date}.\n\nPlease review: https://www.autoreconflow.com"
                                html_body = build_html_email(
                                    "Action Required: Reconciliation Awaiting Approval",
                                    "A reconciliation workpaper has been prepared and submitted for supervisory sign-off under SOX 404.",
                                    {
                                        "Preparer": f"{prep_name} ({prep_id})",
                                        "Submitted At": prep_date,
                                        "Matched Items": f"{len(matched_pd):,} records",
                                        "Exceptions to Review": f"{len(un_a_pd) + len(un_b_pd) + len(fuzzy_pd):,} items"
                                    }
                                )
                                ok, info = send_workflow_email(smtp_cfg, rev_target_email, "ACTION REQUIRED: Reconciliation Sign-Off Pending", plain_body, html_body)
                                st.rerun()
                    else:
                        st.success(f"🔒 Preparer Task Completed & Locked by {st.session_state.get('saved_prep_name', prep_name)}")

                # STAGE 2: REVIEWER PANEL
                with so_col2:
                    if not is_prep_done:
                        st.markdown("##### 🛡️ Stage 2: Reviewer (Supervisory Oversight)")
                        st.info("⏳ Reviewer stage is locked. Awaiting Preparer completion on the left.")
                        rev_name = "Finance Controller"
                        rev_title = "VP - Financial Reporting & Controls"
                        rev_date = now_str
                        rev_verdict = "Under Remediation"
                        wp_ref = "WP-SOX-REC-2026-Q4"
                        audit_notes = "Pending supervisory sign-off."
                    else:
                        st.markdown("##### 🛡️ Stage 2: Reviewer (Active for Evaluation)")
                        st.success("🟢 Preparer completed task. Ready for Supervisory Evaluation.")
                        rev_name = st.text_input("Reviewer Full Name", value="Finance Controller", key="so_rev_name", disabled=is_rev_approved)
                        rev_title = st.text_input("Role / Designation", value="VP - Financial Reporting & Controls", key="so_rev_title", disabled=is_rev_approved)
                        rev_date = st.text_input("Review Date & Time (UTC)", value=st.session_state.get("saved_rev_date", now_str), key="so_rev_date", disabled=is_rev_approved)
                        rev_verdict = st.selectbox("Supervisory Audit Verdict", ["Approved / Cleared", "Approved with Exceptions Noted", "Under Remediation", "Rejected"], key="so_rev_verdict", disabled=is_rev_approved)
                        wp_ref = st.text_input("Audit Workpaper Reference Code", value="WP-SOX-REC-2026-Q4", key="so_wp_ref", disabled=is_rev_approved)
                        audit_notes = st.text_area("General Supervisory Notes", value="All deterministic matches verified. Exceptions evaluated under tolerable limit.", key="so_notes", disabled=is_rev_approved)

                        if not is_rev_approved and query_state["status"] not in ["PENDING_CLARIFICATION"]:
                            if st.button("🛡️ Approve & Finalize Supervisory Sign-Off (No Queries)", type="primary", use_container_width=True, key="btn_rev_approve_direct"):
                                st.session_state["reviewer_approved"] = True
                                st.session_state["saved_rev_name"] = rev_name
                                st.session_state["saved_rev_title"] = rev_title
                                st.session_state["saved_rev_date"] = rev_date
                                st.session_state["saved_rev_verdict"] = rev_verdict
                                st.session_state["saved_wp_ref"] = wp_ref
                                st.session_state["saved_audit_notes"] = audit_notes

                                prep_mail = st.session_state.get("saved_prep_email", "analyst@company.com")
                                html_approved = build_html_email(
                                    "Reconciliation Workpaper Approved & Sealed",
                                    "The supervisory review cycle has been formally cleared. Final audit artifacts are unlocked.",
                                    {"Reviewer": f"{rev_name} ({rev_title})", "Verdict": rev_verdict, "Approved Timestamp": rev_date}
                                )
                                send_workflow_email(smtp_cfg, prep_mail, f"COMPLETED: Reconciliation Workpaper [{wp_ref}] Approved", "Workpaper formally cleared.", html_approved)
                                st.rerun()
                        elif is_rev_approved:
                            st.success(f"✅ Supervisory Sign-Off Formally Approved & Sealed ({st.session_state.get('saved_rev_verdict', rev_verdict)})")

            # --- LINE-ITEM QUERY & CLARIFICATION THREAD ---
            if is_prep_done and not is_rev_approved:
                st.markdown("#### 💬 Supervisory Exception Query & Clarification Thread")
                disputed_items = []
                for col_id_candidate in ["Invoice_ID", "Transaction_Ref", "col_0"]:
                    if col_id_candidate in un_b_pd.columns: disputed_items.extend(un_b_pd[col_id_candidate].astype(str).tolist())
                    if col_id_candidate in fuzzy_pd.columns: disputed_items.extend(fuzzy_pd[col_id_candidate].astype(str).tolist())
                    if col_id_candidate in un_a_pd.columns: disputed_items.extend(un_a_pd[col_id_candidate].astype(str).tolist())
                disputed_items = sorted(list(set([x for x in disputed_items if x])))

                if query_state["status"] in ["NONE", "RESUBMITTED"]:
                    with st.container():
                        st.info("💡 Reviewer Option: Raise a specific transaction query to send back to the Preparer.")
                        sel_item = st.selectbox("Select Line Item Requiring Clarification:", options=disputed_items, key="sel_query_item")
                        q_text = st.text_area("Reviewer Query / Required Explanation:", placeholder="e.g., Variance of 50.00 observed. Please provide bank swift advice.", key="rev_query_input")
                        if st.button("❓ Send Query Back to Preparer for Clarification", type="secondary", key="btn_send_query"):
                            if q_text.strip():
                                query_state["status"] = "PENDING_CLARIFICATION"
                                query_state["item_id"] = sel_item
                                query_state["reviewer_query"] = q_text
                                query_state["reviewer_query_time"] = datetime.datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S UTC")
                                target_prep_mail = st.session_state.get("saved_prep_email", "analyst@company.com")
                                html_query = build_html_email(
                                    f"Clarification Required: Line Item [{sel_item}]",
                                    "The Reviewer has raised an audit observation on this line item.",
                                    {"Item Queried": sel_item, "Reviewer Note": q_text, "Queried At": query_state["reviewer_query_time"]}
                                )
                                send_workflow_email(smtp_cfg, target_prep_mail, f"ACTION REQUIRED: Reviewer Clarification Query [{sel_item}]", q_text, html_query)
                                st.rerun()

                elif query_state["status"] == "PENDING_CLARIFICATION":
                    st.warning(f"⚠️ Action Required by Preparer: Reviewer queried item [{query_state['item_id']}]")
                    st.markdown(f"**Reviewer Observation ({query_state['reviewer_query_time']}):**")
                    st.markdown(f"> *\"{query_state['reviewer_query']}\"*")
                    prep_reply = st.text_area("Preparer Explanation / Response:", key="prep_reply_input")
                    if st.button("📤 Submit Explanation Back to Reviewer", type="primary", key="btn_submit_reply"):
                        if prep_reply.strip():
                            query_state["status"] = "RESUBMITTED"
                            query_state["preparer_reply"] = prep_reply
                            query_state["preparer_reply_time"] = datetime.datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S UTC")
                            target_rev_mail = st.session_state.get("saved_rev_target_email", "controller@company.com")
                            html_reply = build_html_email(
                                f"Clarification Submitted for [{query_state['item_id']}]",
                                "The Preparer has responded to your observation.",
                                {"Item ID": query_state["item_id"], "Preparer Explanation": prep_reply, "Resubmitted At": query_state["preparer_reply_time"]}
                            )
                            send_workflow_email(smtp_cfg, target_rev_mail, f"EXPLANATION SUBMITTED: Item [{query_state['item_id']}]", prep_reply, html_reply)
                            st.rerun()

                elif query_state["status"] == "RESUBMITTED":
                    st.success(f"📥 Preparer Explanation Received for [{query_state['item_id']}].")
                    st.markdown(f"**Reviewer Query:** *\"{query_state['reviewer_query']}\"*")
                    st.markdown(f"**Preparer Explanation ({query_state['preparer_reply_time']}):**")
                    st.markdown(f"> **\"{query_state['preparer_reply']}\"**")
                    if st.button("🛡️ Accept Explanation & Formally Approve Workpaper", type="primary", key="btn_accept_query_apprv"):
                        st.session_state["reviewer_approved"] = True
                        query_state["status"] = "APPROVED"
                        st.session_state["saved_rev_name"] = rev_name
                        st.session_state["saved_rev_title"] = rev_title
                        st.session_state["saved_rev_date"] = datetime.datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S UTC")
                        st.session_state["saved_rev_verdict"] = "Approved with Clarifications Noted"
                        st.session_state["saved_wp_ref"] = wp_ref
                        st.session_state["saved_audit_notes"] = audit_notes
                        st.rerun()

            signoff_payload = {
                "preparer_name": st.session_state.get("saved_prep_name", prep_name),
                "preparer_id": st.session_state.get("saved_prep_id", prep_id),
                "preparer_date": st.session_state.get("saved_prep_date", prep_date),
                "preparer_signed": st.session_state.get("saved_prep_signed", prep_signed),
                "reviewer_name": st.session_state.get("saved_rev_name", rev_name),
                "reviewer_title": st.session_state.get("saved_rev_title", rev_title),
                "reviewer_date": st.session_state.get("saved_rev_date", rev_date),
                "review_verdict": st.session_state.get("saved_rev_verdict", rev_verdict),
                "wp_ref": st.session_state.get("saved_wp_ref", wp_ref),
                "audit_comments": st.session_state.get("saved_audit_notes", audit_notes)
            }
            active_comp_info = {"name": company_name, "address": company_address, "uid": company_uid}
            query_trail_summary = f"Item {query_state['item_id']}: Queried -> Clarified -> Cleared." if query_state.get("status") == "APPROVED" else ""

            # --- OFFICIAL DELIVERABLES HUB ---
            st.markdown("#### 📦 Audit Deliverables & Certification Hub")
            exp_col1, exp_col2 = st.columns(2)

            with exp_col1:
                if is_rev_approved:
                    excel_bytes = generate_consolidated_workbook(
                        matched_pd, fuzzy_pd, un_a_pd, un_b_pd, signoff_payload, 
                        tolerance=res["tolerance"], company_info=active_comp_info, query_obj=query_state
                    )
                    st.download_button(
                        label="📊 Download Final SOX Audit Workpaper (.xlsx)",
                        data=excel_bytes,
                        file_name=f"AutoReconFlow_SOX_Audit_Report_{datetime.datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        use_container_width=True,
                        key="btn_dl_univ_excel"
                    )
                else:
                    st.info("🔒 Excel Workpaper unlocks automatically upon complete Reviewer approval.")

            with exp_col2:
                if is_rev_approved:
                    logo_file_bytes = uploaded_logo.getvalue() if uploaded_logo is not None else None
                    pdf_bytes = generate_pdf_signoff_certificate(
                        len(matched_pd), len(fuzzy_pd), len(un_a_pd), len(un_b_pd), signoff_payload, res["tolerance"], 
                        logo_bytes=logo_file_bytes, company_info=active_comp_info, query_trail=query_trail_summary
                    )
                    st.download_button(
                        label="📜 Download Official Signed PDF Audit Certificate",
                        data=pdf_bytes,
                        file_name=f"AutoReconFlow_SignOff_Certificate_{datetime.datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.pdf",
                        mime="application/pdf",
                        use_container_width=True,
                        key="btn_dl_univ_pdf"
                    )
                else:
                    st.info("🔒 Signed PDF Certificate unlocks automatically once the Reviewer approves.")

            # Four Granular Audit Result Tabs
            tab1, tab2, tab3, tab4 = st.tabs([
                "✅ Tier-1 Deterministic Matches",
                "🔍 Tier-2 Discrepancy & Threshold Review",
                "⚠️ Exceptions: Source A (Ledger Unmatched)",
                "⚠️ Exceptions: Source B (External Unmatched)"
            ])
            with tab1:
                if len(matched_pd) > 0:
                    st.dataframe(matched_pd, use_container_width=True)
                    st.download_button(
                        label="📥 Download Deterministic Matches (CSV)",
                        data=res["matched_pl"].write_csv(),
                        file_name="deterministic_matches.csv",
                        mime="text/csv",
                        key="btn_matched_csv"
                    )
                else:
                    st.info("No exact matches identified under current settings.")

            with tab2:
                if len(fuzzy_pd) > 0:
                    st.markdown("**Discrepancies Requiring Review (Levenshtein / Jaro-Winkler Variance)**")
                    st.dataframe(fuzzy_pd, use_container_width=True)
                    st.download_button(
                        label="📥 Download Discrepancy Candidates (CSV)",
                        data=res["fuzzy_pl"].write_csv(),
                        file_name="fuzzy_review_candidates.csv",
                        mime="text/csv",
                        key="btn_fuzzy_csv"
                    )
                else:
                    st.info("Zero fuzzy discrepancies detected.")

            with tab3:
                if len(un_a_pd) > 0:
                    st.markdown("**Source A Entries Missing in Source B (Unmatched Subledger Items)**")
                    st.dataframe(un_a_pd, use_container_width=True)
                    st.download_button(
                        label="📥 Download Source A Exceptions (CSV)",
                        data=res["un_a_pl"].write_csv(),
                        file_name="unmatched_ledger_exceptions.csv",
                        mime="text/csv",
                        key="btn_un_a_csv"
                    )
                else:
                    st.success("All Source A items accounted for (0 open variances).")

            with tab4:
                if len(un_b_pd) > 0:
                    st.markdown("**Source B Entries Missing in Source A (Timing & External Exceptions)**")
                    st.dataframe(un_b_pd, use_container_width=True)
                    st.download_button(
                        label="📥 Download Source B Exceptions (CSV)",
                        data=res["un_b_pl"].write_csv(),
                        file_name="unmatched_external_exceptions.csv",
                        mime="text/csv",
                        key="btn_un_b_csv"
                    )
                else:
                    st.success("All Source B items accounted for (0 open variances).")
